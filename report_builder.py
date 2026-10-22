"""Genera el 'Production report - DEVESA' (formato USA Kosher) a partir de la bajada de stock.

Toma el modelo de Excel tal cual (modelo/modelo.xlsx) y reemplaza únicamente los datos de la hoja
visible, editando el XML directamente para conservar estilos, anchos, hojas ocultas y fórmulas.

La posición de cada elemento del modelo (encabezados, filas de datos, TOTAL, Spiced, No Spiced,
% Spiced) se detecta leyendo el propio modelo, por lo que no hay posiciones fijas en el código.
"""
import copy
import csv
import datetime as dt
import io
import re
import unicodedata
import zipfile

import openpyxl
import pandas as pd
from lxml import etree
from openpyxl.utils import column_index_from_string, get_column_letter

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
RNS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PNS = "http://schemas.openxmlformats.org/package/2006/relationships"
CTNS = "http://schemas.openxmlformats.org/package/2006/content-types"
q = lambda t: "{%s}%s" % (NS, t)

DESTINO_DEFECTO = "KOSHER USA"
FACTOR_ESCALA = 100000

REQUERIDAS = ["Codigo", "Destino", "MercDescripcion", "Unidad", "Peso", "PesoBruto", "Fecha", "FechaFaena"]


# ------------------------------------------------------------------ detección del layout del modelo
def _norm(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().lower().strip()
    return re.sub(r"[^a-z0-9% ]", "", s)


_ENCABEZADOS = [   # (clave, condición sobre el texto normalizado del encabezado)
    ("barcode", lambda t: t == "barcode"),
    ("producer", lambda t: t == "producer"),
    ("fecha_prod", lambda t: t.startswith("fecha de produc")),
    ("desc", lambda t: "mercdescripcion" in t.replace(" ", "") or t == "descripcion"),
    ("codigo", lambda t: t == "codigo"),
    ("faena", lambda t: "faena" in t),
    ("cajas", lambda t: "unidad" in t),
    ("bruto", lambda t: "pesobruto" in t.replace(" ", "")),
    ("neto", lambda t: "peso" in t and "bruto" not in t),
]
_OBLIGATORIAS = ["fecha_prod", "desc", "codigo", "faena", "cajas", "bruto", "neto"]


def _ruta_hoja(zf, titulo):
    wb = etree.fromstring(zf.read("xl/workbook.xml"))
    rels = etree.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    targets = {r.get("Id"): r.get("Target") for r in rels}
    for sh in wb.find(q("sheets")):
        if sh.get("name") == titulo:
            t = targets[sh.get("{%s}id" % RNS)]
            return t.lstrip("/") if t.startswith("/") else "xl/" + t
    raise ValueError("No se encontró la hoja en el modelo")


def detectar_layout(modelo_path):
    wb = openpyxl.load_workbook(modelo_path)
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        # fila de encabezados: la que más encabezados reconocidos tiene
        mejor, hdr, cols = 0, None, None
        for r in range(1, min(ws.max_row, 30) + 1):
            found = {}
            for c in range(1, ws.max_column + 1):
                v = ws.cell(r, c).value
                if isinstance(v, str):
                    t = _norm(v)
                    for key, fn in _ENCABEZADOS:
                        if key not in found and fn(t):
                            found[key] = c
                            break
            if len(found) > mejor:
                mejor, hdr, cols = len(found), r, found
        if hdr is None or not all(k in cols for k in _OBLIGATORIAS):
            continue
        # filas de totales por etiqueta
        T = lab = None
        for r in range(hdr + 1, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                v = ws.cell(r, c).value
                if isinstance(v, str) and _norm(v) == "total":
                    T, lab = r, c
                    break
            if T:
                break
        if not T:
            continue
        filas = {}
        for r in range(T + 1, ws.max_row + 1):
            v = ws.cell(r, lab).value
            if isinstance(v, str):
                t = _norm(v)
                if t in ("spiced", "no spiced", "% spiced"):
                    filas.setdefault(t, r)
        # fin del rango de datos según la fórmula SUM del total
        end = T - 1
        f = ws.cell(T, cols["cajas"]).value
        if isinstance(f, str):
            m = re.search(r"SUM\(\$?[A-Z]+\$?(\d+):\$?[A-Z]+\$?(\d+)\)", f.upper())
            if m:
                end = int(m.group(2))
        dc = cols["desc"]
        tpl = hdr + 1
        for r in range(hdr + 1, end + 1):
            if ws.cell(r, dc)._style == ws.cell(hdr + 1, dc)._style:
                tpl = r
        prod = ws.cell(hdr + 1, cols["producer"]).value if "producer" in cols else None
        with zipfile.ZipFile(modelo_path) as zf:
            xml = _ruta_hoja(zf, ws.title)
        return {"xml": xml, "hdr": hdr, "start": hdr + 1, "end": end, "tpl": tpl, "T": T, "cols": cols,
                "SP": filas.get("spiced"), "NSP": filas.get("no spiced"), "PCT": filas.get("% spiced"),
                "producer": str(prod) if prod else "AZUL"}
    raise ValueError("No se pudo reconocer la estructura de modelo/modelo.xlsx "
                     "(se esperan los encabezados de columna y la fila TOTAL).")


# ------------------------------------------------------------------ lectura de datos
def cargar_barcodes(path):
    out = {}
    with open(path, encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            cod, bc = (row.get("codigo") or "").strip().upper(), (row.get("barcode") or "").strip()
            if cod and bc:
                out[cod] = bc
    return out


def parsear_barcodes_extra(texto):
    out = {}
    for linea in (texto or "").splitlines():
        partes = [p.strip() for p in re.split(r"[,;\t]", linea) if p.strip()]
        if len(partes) >= 2:
            out[partes[0].upper()] = partes[1]
    return out


def normalizar(df):
    """Unifica los nombres de columna de las distintas bajadas."""
    df = df.copy()
    df.columns = [str(c).strip() for c in df.columns]
    ren = {}
    if "cajas" in df.columns:
        ren["cajas"] = "Unidad"
    if "kg" in df.columns:
        ren["kg"] = "Peso"
    if "kg bruto" in df.columns:
        ren["kg bruto"] = "PesoBruto"
    df = df.rename(columns=ren)
    faltan = [c for c in REQUERIDAS if c not in df.columns]
    if faltan:
        raise ValueError("Al archivo le faltan las columnas: " + ", ".join(faltan))
    df["Fecha"] = pd.to_datetime(df["Fecha"]).dt.normalize()
    df["FechaFaena"] = pd.to_datetime(df["FechaFaena"]).dt.normalize()
    for c in ["Unidad", "Peso", "PesoBruto"]:
        df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0)
    df["Codigo"] = df["Codigo"].astype(str).str.strip().str.upper()
    return df


# ------------------------------------------------------------------ armado de cada día
def _serial(d):
    return (pd.Timestamp(d).to_pydatetime() - dt.datetime(1899, 12, 30)).days


def agrupar(df_dia, barcodes):
    g = df_dia.groupby("Codigo").agg(
        desc=("MercDescripcion", lambda s: s.dropna().astype(str).iloc[0]),
        faena=("FechaFaena", "max"),
        cajas=("Unidad", "sum"),
        bruto=("PesoBruto", "sum"),
        neto=("Peso", "sum"),
    ).reset_index()
    g["barcode"] = g["Codigo"].map(barcodes)
    g["_orden"] = g["barcode"].fillna("~")            # sin barcode, al final
    g = g.sort_values(["_orden", "Codigo"]).reset_index(drop=True)
    g["cajas"] = g["cajas"].round(0)
    g["bruto"] = g["bruto"].round(3)
    g["neto"] = g["neto"].round(3)
    return g


def _desplazar_refs(texto, T_old, end_old, k):
    """Desplaza referencias A1 de una fórmula: filas >= T_old suben k, y el fin del rango de datos también."""
    partes = re.split(r'("[^"]*")', texto)
    def rep(m):
        fila = int(m.group(2))
        if fila >= T_old or fila == end_old:
            fila += k
        return m.group(1) + str(fila)
    return "".join(p if p.startswith('"') else re.sub(r"(\$?[A-Z]{1,3}\$?)(\d+)", rep, p) for p in partes)


def _desplazar_sqref(texto, T_old, k):
    return re.sub(r"([A-Z]{1,3})(\d+)", lambda m: m.group(1) + str(int(m.group(2)) + (k if int(m.group(2)) >= T_old else 0)), texto)


def _workbook_bytes(g, dia, modelo_path, L):
    n = len(g)
    start, end, T, cols, tpl = L["start"], L["end"], L["T"], L["cols"], L["tpl"]
    capacidad = end - start + 1                       # filas dentro del rango de la fórmula de total
    nativas = tpl - start + 1                         # filas con formato de datos
    k = max(0, n - capacidad)                         # filas extra a insertar
    L_ = {key: get_column_letter(c) for key, c in cols.items()}
    colnum = lambda ref: re.match(r"[A-Z]+", ref).group(0)

    z = zipfile.ZipFile(modelo_path)
    infos = [i for i in z.infolist() if i.filename != "xl/calcChain.xml"]
    files = {i.filename: z.read(i.filename) for i in infos}

    # el calcChain se elimina (Excel lo reconstruye solo); evita referencias a celdas movidas
    ct = etree.fromstring(files["[Content_Types].xml"])
    for o in list(ct):
        if o.get("PartName") == "/xl/calcChain.xml":
            ct.remove(o)
    files["[Content_Types].xml"] = etree.tostring(ct, xml_declaration=True, encoding="UTF-8", standalone=True)
    rel = etree.fromstring(files["xl/_rels/workbook.xml.rels"])
    for r in list(rel):
        if r.get("Target", "").endswith("calcChain.xml"):
            rel.remove(r)
    files["xl/_rels/workbook.xml.rels"] = etree.tostring(rel, xml_declaration=True, encoding="UTF-8", standalone=True)

    sst = etree.fromstring(files["xl/sharedStrings.xml"]) if "xl/sharedStrings.xml" in files else None
    idx = {}
    if sst is not None:
        for i, si in enumerate(sst.findall(q("si"))):
            t = si.find(q("t"))
            if t is not None and len(si) == 1:
                idx.setdefault(t.text, i)
    cnt = [len(sst.findall(q("si"))) if sst is not None else 0]

    def sid(text):
        if text in idx:
            return idx[text]
        si = etree.SubElement(sst, q("si"))
        t = etree.SubElement(si, q("t"))
        t.text = text
        if text != text.strip():
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        idx[text] = cnt[0]
        cnt[0] += 1
        return idx[text]

    sh = etree.fromstring(files[L["xml"]])
    sd = sh.find(q("sheetData"))

    B = end + 1                                        # a partir de esta fila todo baja k filas
    if k:
        for r in sd.findall(q("row")):
            rn = int(r.get("r"))
            if rn >= B:
                r.set("r", str(rn + k))
                for c in r.findall(q("c")):
                    c.set("r", colnum(c.get("r")) + str(rn + k))
        for f in sh.iter(q("f")):
            if f.text:
                f.text = _desplazar_refs(f.text, B, end, k)
        for e in sh.iter(q("mergeCell")):
            e.set("ref", _desplazar_sqref(e.get("ref"), B, k))
        for e in sh.iter(q("conditionalFormatting")):
            e.set("sqref", _desplazar_sqref(e.get("sqref"), B, k))
        for e in sh.iter(q("dataValidation")):
            e.set("sqref", _desplazar_sqref(e.get("sqref"), B, k))
    rows = {int(r.get("r")): r for r in sd.findall(q("row"))}
    rt = rows[tpl]
    tpl_cells = {colnum(c.get("r")): c for c in rt.findall(q("c"))}

    def fila(rn):
        """Devuelve la fila rn; si no existe en el XML (filas vacías), la crea con el formato de la fila modelo."""
        if rn in rows:
            return rows[rn]
        new = etree.Element(q("row"))
        new.set("r", str(rn))
        for a_ in ("spans", "ht", "customHeight"):
            if rt.get(a_) is not None:
                new.set(a_, rt.get(a_))
        pos = [i for i, r in enumerate(sd.findall(q("row"))) if int(r.get("r")) > rn]
        sd.insert(pos[0], new) if pos else sd.append(new)
        rows[rn] = new
        return new

    def celda(rn, col):
        row = fila(rn)
        ref = f"{col}{rn}"
        for c in row.findall(q("c")):
            if c.get("r") == ref:
                return c
        new = etree.Element(q("c"))
        new.set("r", ref)
        if col in tpl_cells and tpl_cells[col].get("s") is not None:
            new.set("s", tpl_cells[col].get("s"))
        ci = column_index_from_string(col)
        pos = [i for i, c in enumerate(row.findall(q("c"))) if column_index_from_string(colnum(c.get("r"))) > ci]
        if pos:
            row.insert(list(row).index(row.findall(q("c"))[pos[0]]), new)
        else:
            row.append(new)
        return new

    if n > nativas:                                    # filas que hacen falta con el formato de datos
        for rn in range(tpl + 1, start + n):
            row = fila(rn)
            if rn <= end:                              # fila existente del rango (separadora): se le da formato de datos
                for col, c0 in tpl_cells.items():
                    c = celda(rn, col)
                    c.attrib.pop("s", None)
                    if c0.get("s") is not None:
                        c.set("s", c0.get("s"))
                for c in row.findall(q("c")):
                    if colnum(c.get("r")) not in tpl_cells:
                        c.attrib.pop("s", None)
                if rt.get("ht") is not None:
                    row.set("ht", rt.get("ht"))
            else:                                      # fila nueva insertada
                for col in tpl_cells:
                    celda(rn, col)

    def setcell(r, col, val, kind):
        c = celda(r, col)
        for ch in list(c):
            c.remove(ch)
        c.attrib.pop("t", None)
        if val is None:
            return
        v = etree.SubElement(c, q("v"))
        if kind == "s" and sst is None:              # modelo sin tabla de textos compartidos
            c.remove(v)
            c.set("t", "inlineStr")
            t = etree.SubElement(etree.SubElement(c, q("is")), q("t"))
            t.text = val
        elif kind == "s":
            c.set("t", "s")
            v.text = str(sid(val))
        else:
            v.text = repr(val) if isinstance(val, float) else str(val)

    S = (pd.Timestamp(dia).to_pydatetime() - dt.datetime(1899, 12, 30)).days
    serial = lambda d: (pd.Timestamp(d).to_pydatetime() - dt.datetime(1899, 12, 30)).days
    for i in range(start, start + max(nativas, n)):
        j = i - start
        if j < n:
            x = g.iloc[j]
            if "barcode" in L_:
                setcell(i, L_["barcode"], x["barcode"] if isinstance(x["barcode"], str) else None, "s")
            if "producer" in L_:
                setcell(i, L_["producer"], L["producer"], "s")
            setcell(i, L_["fecha_prod"], S, "n")
            setcell(i, L_["desc"], x["desc"], "s")
            setcell(i, L_["codigo"], x["Codigo"], "s")
            setcell(i, L_["faena"], serial(x["faena"]), "n")
            setcell(i, L_["cajas"], int(x["cajas"]), "n")
            setcell(i, L_["bruto"], float(x["bruto"]), "n")
            setcell(i, L_["neto"], float(x["neto"]), "n")
        else:
            for col in L_.values():
                setcell(i, col, None, "n")

    # valores en caché de las filas de totales (las fórmulas son las del modelo, con el rango ya ajustado)
    sp = g[g["desc"].str.upper().str.contains("SPICED")]
    def val(key, df_):
        return int(df_[{"cajas": "cajas", "bruto": "bruto", "neto": "neto"}[key]].sum()) if key == "cajas" \
            else round(float(df_[{"bruto": "bruto", "neto": "neto"}[key]].sum()), 3)
    roles = {}
    for key in ("cajas", "bruto", "neto"):
        tot, spv = val(key, g), val(key, sp)
        roles[(T, key)] = tot
        if L["SP"]: roles[(L["SP"], key)] = spv
        if L["NSP"]: roles[(L["NSP"], key)] = tot - spv if key == "cajas" else round(tot - spv, 3)
        if L["PCT"]: roles[(L["PCT"], key)] = (spv / tot) if tot else 0
    for (fila, key), v in roles.items():
        ref = f"{L_[key]}{fila + k}"
        for c in rows.get(fila + k, etree.Element("x")).findall(q("c")):
            if c.get("r") == ref and c.find(q("f")) is not None:
                vv = c.find(q("v"))
                if vv is None:
                    vv = etree.SubElement(c, q("v"))
                vv.text = repr(v)

    pane = sh.find(q("sheetViews")).find(q("sheetView")).find(q("pane"))
    if pane is not None:
        pane.set("topLeftCell", f"A{start}")
    af_old = None
    if k:
        dim = sh.find(q("dimension"))
        if dim is not None:
            m = re.match(r"(.*?)(\d+)$", dim.get("ref"))
            dim.set("ref", m.group(1) + str(int(m.group(2)) + k))
        af = sh.find(q("autoFilter"))
        if af is not None:
            af_old = af.get("ref")
            m = re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)$", af_old)
            af.set("ref", f"{m.group(1)}{m.group(2)}:{m.group(3)}{int(m.group(4)) + k}")
    wbx = etree.fromstring(files["xl/workbook.xml"])
    if af_old:
        m = re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)$", af_old)
        old_abs = f"${m.group(1)}${m.group(2)}:${m.group(3)}${m.group(4)}"
        new_abs = f"${m.group(1)}${m.group(2)}:${m.group(3)}${int(m.group(4)) + k}"
        for dn in wbx.iter(q("definedName")):
            if dn.text and dn.text.endswith(old_abs) and "_FilterDatabase" in (dn.get("name") or ""):
                dn.text = dn.text[: -len(old_abs)] + new_abs
    calc = wbx.find(q("calcPr"))
    if calc is None:
        calc = etree.SubElement(wbx, q("calcPr"))
    calc.set("fullCalcOnLoad", "1")                    # Excel recalcula al abrir
    files["xl/workbook.xml"] = etree.tostring(wbx, xml_declaration=True, encoding="UTF-8", standalone=True)

    files[L["xml"]] = etree.tostring(sh, xml_declaration=True, encoding="UTF-8", standalone=True)
    if sst is not None:
        sst.set("uniqueCount", str(cnt[0]))
        sst.set("count", str(int(sst.get("count")) + n * 3))
        files["xl/sharedStrings.xml"] = etree.tostring(sst, xml_declaration=True, encoding="UTF-8", standalone=True)

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zo:
        for i in infos:
            zo.writestr(i.filename, files[i.filename])
    return buf.getvalue()


def generar_reportes(df, barcodes, modelo_path, destino=DESTINO_DEFECTO):
    """Devuelve (reportes, avisos). Un reporte por cada fecha con producción del destino."""
    df = normalizar(df)
    avisos = []
    kdf = df[df["Destino"].astype(str).str.strip().str.upper() == destino.strip().upper()].copy()
    if kdf.empty:
        destinos = ", ".join(sorted(df["Destino"].dropna().astype(str).unique())[:15])
        raise ValueError(f"No hay registros con destino «{destino}». Destinos en el archivo: {destinos}")

    # Algunas bajadas traen los pesos multiplicados por 100.000: se detecta por kg/caja
    cajas, neto = kdf["Unidad"].sum(), kdf["Peso"].sum()
    if cajas > 0 and neto / cajas > 1000:
        kdf["Peso"] /= FACTOR_ESCALA
        kdf["PesoBruto"] /= FACTOR_ESCALA
        avisos.append(f"Los pesos venían multiplicados por {FACTOR_ESCALA:,}; se dividieron automáticamente."
                      .replace(",", "."))

    L = detectar_layout(modelo_path)
    reportes = []
    for dia in sorted(kdf["Fecha"].unique()):
        dia = pd.Timestamp(dia)
        g = agrupar(kdf[kdf["Fecha"] == dia], barcodes)
        data = _workbook_bytes(g, dia, modelo_path, L)
        sp = g[g["desc"].str.upper().str.contains("SPICED")]
        tot_c = float(g["cajas"].sum())
        sin_bc = g[g["barcode"].isna()][["Codigo", "desc", "cajas"]].to_dict("records")
        reportes.append({
            "fecha": dia.strftime("%d/%m/%Y"),
            "archivo": f"Production_report_-_DEVESA_{dia:%d-%m}_-{dia:%Y}_USA_-.xlsx",
            "bytes": data,
            "productos": len(g),
            "cajas": int(tot_c),
            "bruto": round(float(g["bruto"].sum()), 3),
            "neto": round(float(g["neto"].sum()), 3),
            "pct_spiced": (float(sp["cajas"].sum()) / tot_c * 100) if tot_c else 0.0,
            "sin_barcode": sin_bc,
        })
    return reportes, avisos
