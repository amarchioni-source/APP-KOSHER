"""Genera el 'Production report - DEVESA' de Israel (Kosher / NETO) a partir de la bajada de stock.

Una fila por cartón, con el código de caja (CODBAR CAJA) y la fecha de vencimiento (BEST BEFORE DATE).
Usa modelo/modelo_israel.xlsx como plantilla (edición directa del XML, para conservar estilos y anchos).
"""
import csv
import datetime as dt
import io
import re
import zipfile

import openpyxl
import pandas as pd
from lxml import etree
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
RNS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
q = lambda t: "{%s}%s" % (NS, t)

PALABRAS = ("NETO", "OZUL", "T BONE VEAL", "GRASA SALADA")
ESTADOS_OK = {"E", "PA", "D", "104"}
FACTOR = 100000
COLS = "ABCDEFGHIJKLM"
PLAZOS = ("Plazos habituales según el lote: entre 81 y 88 días desde la producción en los cortes, 365 días en "
          "timus, mollejas y sesos, y 456 o 457 días en lengua y tendones. No se estimó ninguna fecha.")

ALIAS = {   # clave -> nombres posibles de columna (normalizados: minúsculas, sin espacios ni símbolos)
    "codigo": ["codigo"],
    "desc": ["mercdescripcion", "descripcion"],
    "estado": ["estado"],
    "codbars": ["codbars", "codbar", "codbarcaja"],
    "faena": ["fechafaena"],
    "prod": ["fechaproduccion", "fecha"],
    "venc": ["fechavencimiento", "vencimiento"],
    "units": ["units", "cantidad", "unidades"],
    "cartons": ["cartons", "unidad", "cajas"],
    "neto": ["pesoneto", "peso", "kg"],
    "bruto": ["pesobruto", "kgbruto"],
}
_NOMBRES = {"codigo": "Codigo", "desc": "MercDescripcion", "estado": "Estado", "codbars": "CodBarS",
            "faena": "FechaFaena", "prod": "FechaProduccion", "units": "units/Cantidad",
            "neto": "Peso (neto)", "bruto": "PesoBruto"}


# ------------------------------------------------------------------ tabla de índices
def cargar_indices(path):
    out = {}
    with open(path, encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f):
            cod = (r.get("codigo") or "").strip().upper()
            if cod:
                out[cod] = {"indice": (r.get("indice") or "").strip(), "barcode": (r.get("barcode") or "").strip(),
                            "desc": (r.get("descripcion") or "").strip(),
                            "inferido": (r.get("inferido") or "").strip().upper() in ("SI", "SÍ", "1", "X", "INFERIDO")}
    return out


def parsear_indices_extra(texto):
    """Líneas 'CODIGO,indice,barcode' (para una corrida puntual)."""
    out = {}
    for linea in (texto or "").splitlines():
        p = [x.strip() for x in re.split(r"[,;\t|]", linea) if x.strip() != ""]
        if len(p) >= 3:
            out[p[0].upper()] = {"indice": p[1], "barcode": p[2], "desc": "", "inferido": False}
    return out


# ------------------------------------------------------------------ lectura y normalización
def _k(c):
    return re.sub(r"[^a-z0-9]", "", str(c).lower())


def _estado(v):
    if pd.isna(v):
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).strip().upper()


def _codbar(v):
    if pd.isna(v):
        return ""
    if isinstance(v, (int, float)):
        return str(int(v))
    s = str(v).strip()
    return re.sub(r"\.0$", "", s) if re.fullmatch(r"\d+\.0", s) else s


def _fecha(s):
    if pd.api.types.is_datetime64_any_dtype(s):
        return s.dt.normalize()
    return pd.to_datetime(s, dayfirst=True, errors="coerce").dt.normalize()


def normalizar_stock(df):
    """Devuelve (DataFrame normalizado, avisos). Una fila por cartón."""
    cols = {_k(c): c for c in df.columns}
    m = {key: next((cols[a] for a in alias if a in cols), None) for key, alias in ALIAS.items()}
    faltan = [n for key, n in _NOMBRES.items() if not m[key]]
    if faltan:
        raise ValueError("No parece una bajada por cartón (formato Israel). Faltan las columnas: " + ", ".join(faltan))
    num = lambda c: pd.to_numeric(df[c], errors="coerce").fillna(0)
    out = pd.DataFrame({
        "Codigo": df[m["codigo"]].astype(str).str.strip().str.upper(),
        "desc": df[m["desc"]].astype(str).str.strip(),
        "estado": df[m["estado"]].map(_estado),
        "codbars": df[m["codbars"]].map(_codbar),
        "faena": _fecha(df[m["faena"]]),
        "prod": _fecha(df[m["prod"]]),
        "venc": _fecha(df[m["venc"]]) if m["venc"] else pd.Series(pd.NaT, index=df.index),
        "units": num(m["units"]),
        "cartons": num(m["cartons"]) if m["cartons"] else 1,
        "neto": num(m["neto"]),
        "bruto": num(m["bruto"]),
    })
    out = out[~out["Codigo"].isin(["", "NAN", "NONE"])].copy()
    avisos = []
    if len(out):
        if out["neto"].median() > 1000:
            out["neto"] /= FACTOR
            out["bruto"] /= FACTOR
            avisos.append("Los pesos venían multiplicados por 100.000; se dividieron automáticamente.")
        if out["units"].median() > 5000:
            out["units"] /= FACTOR
            avisos.append("Las unidades venían multiplicadas por 100.000; se dividieron automáticamente.")
        por_caja = out["neto"].sum() / max(out["cartons"].sum(), 1)
        if not 3 <= por_caja <= 60:
            avisos.append(f"Revisar los pesos: el promedio es de {por_caja:,.1f} kg netos por cartón.".replace(",", "."))
    return out, avisos


# ------------------------------------------------------------------ plantilla
def _ruta_hoja(zf, titulo):
    wb = etree.fromstring(zf.read("xl/workbook.xml"))
    rels = etree.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    targets = {r.get("Id"): r.get("Target") for r in rels}
    for sh in wb.find(q("sheets")):
        if sh.get("name") == titulo:
            t = targets[sh.get("{%s}id" % RNS)]
            return t.lstrip("/") if t.startswith("/") else "xl/" + t
    raise ValueError("No se encontró la hoja del modelo de Israel")


def detectar_layout(modelo_path):
    ws = None
    for w in openpyxl.load_workbook(modelo_path).worksheets:
        if w.sheet_state == "visible" and str(w["A2"].value or "").strip().upper() == "PRODUCER":
            ws = w
            break
    if ws is None:
        raise ValueError("No se pudo reconocer modelo/modelo_israel.xlsx (se espera PRODUCER en A2).")
    T = next((r for r in range(ws.max_row, 2, -1) if str(ws.cell(r, 7).value or "").strip().upper() == "TOTAL"), None)
    if not T:
        raise ValueError("El modelo de Israel no tiene fila TOTAL.")
    with zipfile.ZipFile(modelo_path) as zf:
        xml = _ruta_hoja(zf, ws.title)
    return {"xml": xml, "start": 3, "T": T}


def _serial(d):
    return (pd.Timestamp(d).to_pydatetime() - dt.datetime(1899, 12, 30)).days


def _xlsx_dia(g, dia, modelo_path, lay):
    start = lay["start"]
    z = zipfile.ZipFile(modelo_path)
    infos = z.infolist()
    files = {i.filename: z.read(i.filename) for i in infos}
    sh = etree.fromstring(files[lay["xml"]])
    sd = sh.find(q("sheetData"))
    rows = {int(r.get("r")): r for r in sd.findall(q("row"))}
    colof = lambda c: re.match(r"[A-Z]+", c.get("r")).group(0)
    tpl, tot = rows[start], rows[lay["T"]]
    est_d = {colof(c): c.get("s") for c in tpl.findall(q("c"))}
    est_t = {colof(c): c.get("s") for c in tot.findall(q("c"))}
    attr_d = {k: v for k, v in tpl.attrib.items() if k not in ("r", "spans")}
    attr_t = {k: v for k, v in tot.attrib.items() if k not in ("r", "spans")}
    for r in list(sd.findall(q("row"))):
        if int(r.get("r")) >= start:
            sd.remove(r)
    for r in sd.findall(q("row")):               # encabezados: se elimina lo que haya fuera de A:M (ej. el 100000 de N2)
        for c in list(r.findall(q("c"))):
            if colof(c) not in COLS:
                r.remove(c)

    def nuevo_row(rn, attrs):
        row = etree.SubElement(sd, q("row"))
        row.set("r", str(rn))
        for k, v in attrs.items():
            row.set(k, v)
        return row

    def celda(row, rn, col, estilos, kind=None, val=None):
        c = etree.SubElement(row, q("c"))
        c.set("r", f"{col}{rn}")
        if estilos.get(col):
            c.set("s", estilos[col])
        if val is None or val == "":
            return c
        if kind == "t":
            c.set("t", "inlineStr")
            t = etree.SubElement(etree.SubElement(c, q("is")), q("t"))
            t.text = str(val)
        else:
            c.set("t", "n")
            etree.SubElement(c, q("v")).text = str(val)
        return c

    def num(x):
        x = float(x)
        return str(int(x)) if x == int(x) else repr(round(x, 3))

    n = len(g)
    for i, x in enumerate(g.itertuples(index=False)):
        rn = start + i
        row = nuevo_row(rn, attr_d)
        celda(row, rn, "A", est_d, "t", "AZUL")
        celda(row, rn, "B", est_d, "n", _serial(dia))
        celda(row, rn, "C", est_d, "t", x.desc)
        celda(row, rn, "D", est_d, "t", x.indice)
        celda(row, rn, "E", est_d, "t", x.Codigo)
        celda(row, rn, "F", est_d, "t", x.barcode)
        celda(row, rn, "G", est_d, "n", None if pd.isna(x.faena) else _serial(x.faena))
        celda(row, rn, "H", est_d, "n", 1)
        celda(row, rn, "I", est_d, "t", x.codbars)
        celda(row, rn, "J", est_d, "n", num(x.units))
        celda(row, rn, "K", est_d, "n", num(round(x.bruto, 3)))
        celda(row, rn, "L", est_d, "n", num(round(x.neto, 3)))
        celda(row, rn, "M", est_d, "n", None if pd.isna(x.venc) else _serial(x.venc))
    last, T = start + n - 1, start + n
    row = nuevo_row(T, attr_t)
    suma = {"H": n, "J": round(float(g["units"].sum()), 3), "K": round(float(g["bruto"].sum()), 3),
            "L": round(float(g["neto"].sum()), 3)}
    for col in COLS:
        if col == "A":
            celda(row, T, col, est_t, "t", "AZUL")
        elif col == "G":
            celda(row, T, col, est_t, "t", "TOTAL")
        elif col in suma:
            c = celda(row, T, col, est_t)
            c.set("t", "n")
            etree.SubElement(c, q("f")).text = f"SUM({col}{start}:{col}{last})"
            etree.SubElement(c, q("v")).text = num(suma[col])
        else:
            celda(row, T, col, est_t)

    dim = sh.find(q("dimension"))
    if dim is not None:
        dim.set("ref", f"A1:M{T}")
    af = sh.find(q("autoFilter"))
    old_abs = None
    if af is not None:
        m = re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)$", af.get("ref"))
        old_abs = f"${m.group(1)}${m.group(2)}:${m.group(3)}${m.group(4)}"
        new_abs = f"${m.group(1)}${m.group(2)}:${m.group(3)}${last}"
        af.set("ref", f"{m.group(1)}{m.group(2)}:{m.group(3)}{last}")
    sv = sh.find(q("sheetViews")).find(q("sheetView"))
    sv.set("topLeftCell", "A1")
    for sel in sv.findall(q("selection")):
        ref = "A1" if sel.get("pane") == "topLeft" else "A3"
        sel.set("activeCell", ref)
        sel.set("sqref", ref)
    wbx = etree.fromstring(files["xl/workbook.xml"])
    if old_abs:
        for dn in wbx.iter(q("definedName")):
            if dn.text and dn.text.endswith(old_abs) and "_FilterDatabase" in (dn.get("name") or ""):
                dn.text = dn.text[: -len(old_abs)] + new_abs
    calc = wbx.find(q("calcPr"))
    if calc is None:
        calc = etree.Element(q("calcPr"))
        ref_el = wbx.find(q("definedNames"))
        if ref_el is None:
            ref_el = wbx.find(q("sheets"))
        ref_el.addnext(calc)
    calc.set("fullCalcOnLoad", "1")
    files["xl/workbook.xml"] = etree.tostring(wbx, xml_declaration=True, encoding="UTF-8", standalone=True)
    files[lay["xml"]] = etree.tostring(sh, xml_declaration=True, encoding="UTF-8", standalone=True)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zo:
        for i in infos:
            zo.writestr(i.filename, files[i.filename])
    return buf.getvalue()


# ------------------------------------------------------------------ generación
def generar_israel(df, indices, modelo_path):
    """df: stock ya normalizado (una o varias bajadas). Devuelve (reportes, control, avisos)."""
    avisos = []
    d = df.reset_index(drop=True)                # las bajadas combinadas repiten etiquetas de fila
    # cartones repetidos entre archivos (mismo CodBarS)
    dup = (d["codbars"] != "") & d.duplicated("codbars", keep="first")
    if dup.sum():
        d = d[~dup]
        avisos.append(f"{int(dup.sum())} cartones estaban repetidos entre archivos (mismo CodBarS) y se contaron una sola vez.")
    # productos que entran
    entra = d["desc"].str.upper().map(lambda s: any(p in s for p in PALABRAS))
    excl = d[~entra]
    excl_desc = (excl.groupby("desc")["cartons"].sum().sort_values(ascending=False)
                 .rename("cartones").reset_index().to_dict("records"))
    d = d[entra]
    # estados
    s_cart = int(d.loc[d["estado"] == "S", "cartons"].sum())
    otros = d[~d["estado"].isin(ESTADOS_OK | {"S"})]
    estados_otros = (otros.groupby("estado")["cartons"].sum().rename("cartones").reset_index().to_dict("records"))
    d = d[d["estado"].isin(ESTADOS_OK)]
    sin_dia = int(d["prod"].isna().sum())
    if sin_dia:
        avisos.append(f"{sin_dia} cartones no tienen fecha de producción y quedaron fuera.")
        d = d[d["prod"].notna()]
    if d.empty:
        raise ValueError("No quedaron cartones para el informe de Israel (revise descripciones y estados del archivo).")
    if (d["cartons"] != 1).any():
        avisos.append(f"{int((d['cartons'] != 1).sum())} filas no traen exactamente 1 cartón; el reporte muestra 1 cartón por fila.")

    lay = detectar_layout(modelo_path)
    reportes, filas_ctrl = [], []
    for dia in sorted(d["prod"].unique()):
        dia = pd.Timestamp(dia)
        g = d[d["prod"] == dia].copy()
        g["indice"] = g["Codigo"].map(lambda c: indices.get(c, {}).get("indice", ""))
        g["barcode"] = g["Codigo"].map(lambda c: indices.get(c, {}).get("barcode", ""))
        g["_grupo"] = g["Codigo"].str.startswith("FD").map({True: 0, False: 1})
        g["_cb"] = g["codbars"].replace("", "~")
        g["_fa"] = g["faena"].fillna(pd.Timestamp.max.normalize())
        g = g.sort_values(["_grupo", "Codigo", "_fa", "_cb"], kind="stable").reset_index(drop=True)
        data = _xlsx_dia(g, dia, modelo_path, lay)
        faltan = [{"Codigo": c, "desc": x["desc"].iloc[0], "cartones": len(x)}
                  for c, x in g[(g["indice"] == "") | (g["barcode"] == "")].groupby("Codigo")]
        infer = [{"Codigo": c, "desc": x["desc"].iloc[0], "cartones": len(x)}
                 for c, x in g.groupby("Codigo") if indices.get(c, {}).get("inferido")]
        sin_cod = int(((g["codbars"] == "") & (g["estado"] != "104")).sum())
        sin_104 = int(((g["codbars"] == "") & (g["estado"] == "104")).sum())
        rep = {"dia": dia, "fecha": dia.strftime("%d/%m/%Y"),
               "archivo": f"Production report - DEVESA {dia:%d-%m-%Y}.xlsx", "bytes": data,
               "productos": g["Codigo"].nunique(), "cartones": len(g), "unidades": float(g["units"].sum()),
               "bruto": round(float(g["bruto"].sum()), 3), "neto": round(float(g["neto"].sum()), 3),
               "faltan_indice": faltan, "inferidos": infer, "sin_venc": int(g["venc"].isna().sum()),
               "sin_codbar": sin_cod, "sin_codbar_104": sin_104}
        reportes.append(rep)
    fechas = [r["dia"] for r in reportes]
    habiles = pd.bdate_range(min(fechas), max(fechas))
    sin_archivo = [x.strftime("%d/%m/%Y") for x in habiles if x not in set(fechas)]
    control = {"excl_desc": excl_desc, "s_cartones": s_cart, "estados_otros": estados_otros,
               "dias_sin_archivo": sin_archivo}
    return reportes, control, avisos


# ------------------------------------------------------------------ Excel de control
def crear_control(reportes, control):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumen por día"
    ws.append(["Fecha", "Cartones", "Unidades", "Kg bruto", "Kg neto"])
    for r in reportes:
        ws.append([r["dia"].to_pydatetime(), r["cartones"], r["unidades"], r["bruto"], r["neto"]])
    n = len(reportes)
    ws.append(["TOTAL", f"=SUM(B2:B{n + 1})", f"=SUM(C2:C{n + 1})", f"=SUM(D2:D{n + 1})", f"=SUM(E2:E{n + 1})"])
    for c in ws[1] + ws[n + 2]:
        c.font = Font(bold=True)
    for c in ws[1]:
        c.fill = PatternFill("solid", fgColor="E8EEF5")
        c.alignment = Alignment(horizontal="center")
    for row in ws.iter_rows(min_row=2, max_row=n + 2):
        row[0].number_format = "dd/mm/yyyy"
        row[1].number_format = "#,##0"
        row[2].number_format = "#,##0"
        row[3].number_format = "#,##0.000"
        row[4].number_format = "#,##0.000"
    for col, w in zip("ABCDE", (14, 12, 12, 16, 16)):
        ws.column_dimensions[col].width = w

    p = wb.create_sheet("Pendientes")

    def bloque(titulo, cabecera, filas):
        p.append([titulo])
        p.cell(p.max_row, 1).font = Font(bold=True, size=12)
        if filas:
            p.append(cabecera)
            for c in p[p.max_row]:
                c.font = Font(bold=True)
                c.fill = PatternFill("solid", fgColor="E8EEF5")
            for f in filas:
                p.append(f)
        else:
            p.append(["Sin pendientes"])
        p.append([])

    bloque("Códigos sin índice NETO o barcode (celdas vacías en el reporte)", ["Fecha", "Código", "Descripción", "Cartones"],
           [[r["fecha"], f["Codigo"], f["desc"], f["cartones"]] for r in reportes for f in r["faltan_indice"]])
    bloque("Cartones sin fecha de vencimiento (no se estimó)", ["Fecha", "Cartones"],
           [[r["fecha"], r["sin_venc"]] for r in reportes if r["sin_venc"]])
    bloque("Códigos con índice y barcode inferidos (no confirmados)", ["Fecha", "Código", "Descripción", "Cartones"],
           [[r["fecha"], f["Codigo"], f["desc"], f["cartones"]] for r in reportes for f in r["inferidos"]])
    bloque("Líneas excluidas por descripción", ["Descripción", "Cartones"],
           [[e["desc"], e["cartones"]] for e in control["excl_desc"]])
    bloque("Cartones en estado S (con salida) excluidos", ["Cartones"], [[control["s_cartones"]]] if control["s_cartones"] else [])
    bloque("Estados no contemplados (excluidos)", ["Estado", "Cartones"],
           [[e["estado"], e["cartones"]] for e in control["estados_otros"]])
    bloque("Días hábiles sin archivo dentro del período", ["Fecha"], [[x] for x in control["dias_sin_archivo"]])
    for col, w in zip("ABCD", (62, 14, 50, 12)):
        p.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
